import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

P="/mnt/user-data/outputs/Henkel_MA_Financing_Model_teilverkauf_ccy.xlsx"
wb=openpyxl.load_workbook(P)

FONT="Arial"
def F(sz=10,b=False,color="000000",it=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=it)
def fill(c): return PatternFill("solid",fgColor=c)
YELLOW=fill("FFF2CC"); HEADERBG=fill("1F3864"); BANDBG=fill("EEF3FA"); NOFILL=PatternFill()
thin=Side(style="thin",color="BFBFBF"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin); NOBORDER=Border()
BLUE="0000FF"; GREEN="008000"; GREY="555555"; WHITE="FFFFFF"
NUM_MM='#,##0.0;(#,##0.0);"-"'; PCT1='0.0%'; PCT2='0.00%'; EPS4='0.0000'
def put(ws,a,v,font=None,fillc=None,num=None,align=None,wrap=False,border=False):
    c=ws[a]; c.value=v; c.font=font or F()
    if fillc: c.fill=fillc
    if num: c.number_format=num
    if align or wrap: c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if border: c.border=BORDER
    return c

NY=8
MIX="'3. Scenario & Mix'"; DI="'1. Deal Inputs'"; E3="'E3 · Cost & Schedule'"
E4="'E4 · KPIs & Funding'"; E1="'E1 · Rates & Spreads'"
OPTAX=f"{DI}!$B$42"; SHIELD=f"{E4}!$B$96"; SHARES=f"{DI}!$B$44"; ENT=f"{DI}!$B$65"

mix=wb["3. Scenario & Mix"]; hs=wb["Henkel Summary"]

# ==================== TAB 3: add EBIT-share column D to sales table ====================
put(mix,"D85","EBIT share of target (%)",F(8,True),BANDBG,align="center",wrap=True,border=True)
put(mix,"E85","(blank row = no sale; blank D = NES proxy)",F(8,False,GREY,it=True))
for k in range(6):
    put(mix,f"D{86+k}",None,F(9,True,BLUE),YELLOW,PCT1,align="center",border=True)

# ==================== TAB 3: clear obsolete col H (per-ccy margins) + col K (helper) ====================
c=mix["H7"]; c.value=None; c.font=F(8); c.fill=NOFILL; c.border=NOBORDER
for r in range(8,30):
    c=mix[f"H{r}"]; c.value=None; c.fill=NOFILL; c.border=NOBORDER
c=mix["K7"]; c.value=None
for r in range(8,30): mix[f"K{r}"].value=None
mix.column_dimensions["K"].hidden=False

# ==================== TAB 3: update note row 92 ====================
put(mix,"A92","Each row removes that currency's financing leg (amount + its own rate). Col D = sold entity's EBIT share (blank = NES proxy). Same-currency fractions add up (cap 100%). Blank B = 100%.",F(8,False,GREY,it=True))

# ==================== HENKEL SUMMARY: rewrite N96-N101 (EBIT share helper) ====================
for k in range(6):
    tr=86+k; rr=96+k
    # direct EBIT share from Tab3 col D; fallback = NES share
    put(hs,f"N{rr}",
        f'=IFERROR(IF(ISNUMBER({MIX}!$D${tr}),{MIX}!$D${tr},'
        f'INDEX({MIX}!$G$8:$G$29,$M${rr})/SUM({MIX}!$G$8:$G$29)),0)',
        F(7,False,GREY))

# ==================== HENKEL SUMMARY: update labels ====================
put(hs,"A79","Operating-earnings retained factor (EBIT share per sale — direct input col D, or NES proxy)",F(9))

# ==================== HENKEL SUMMARY: update echo table ====================
put(hs,"A82","Per sale (echo): ccy · % leg · year · EBIT share · rate(Y1) · Σ interest shed",F(8,True,GREY))
for k in range(6):
    r=83+k; tr=86+k; rr=96+k
    # A-D unchanged (Sale#, ccy, fraction, year)
    put(hs,f"E{r}",f'=IF({MIX}!$A${tr}="","",$N${rr})',F(8),None,PCT1,align="center")           # EBIT share
    put(hs,f"F{r}",f'=IF({MIX}!$A${tr}="","",$C${rr})',F(8),None,PCT2,align="center")            # rate Y1
    put(hs,f"G{r}",f'=IF({MIX}!$A${tr}="","",SUM(C{104+k}:J{104+k}))',F(8),None,NUM_MM,align="center")  # Σ interest

# ==================== HENKEL SUMMARY: update notes (rows 89-90) ====================
for mr in list(hs.merged_cells.ranges):
    if mr.min_row in (89,90) and mr.max_row in (89,90):
        hs.unmerge_cells(str(mr))
put(hs,"A89","Rows 75–80 aggregate ALL sales. Echo: E = EBIT share (direct from Tab 3 col D, or NES proxy); F = Y1 rate; G = total interest removed over horizon.",F(8,False,GREY,it=True)); hs.merge_cells("A89:J89")
put(hs,"A90","Same-currency rows add up (cap 100%). Operating: D filled → exact EBIT share; blank → NES share (revenue proxy). NFP = retained amount × Group NFP.",F(8,False,GREY,it=True)); hs.merge_cells("A90:J90")

# ==================== HENKEL SUMMARY: update post-sale section (91-94) ====================
put(hs,"A91","► POST-SALE OPERATING & EPS — acquirer (EBIT share from Tab 3 col D; blank = NES proxy)",F(11,True,WHITE),HEADERBG)
for mr in list(hs.merged_cells.ranges):
    if mr.min_row==94 and mr.max_row==94:
        hs.unmerge_cells(str(mr))
put(hs,"A94",
    f'=IF(SUMPRODUCT(--ISNUMBER({MIX}!$D$86:$D$91))>0,'
    f'"Operating EBIT share: direct input (Tab 3 col D). Financing after-tax; WHT scales with remaining interest.",'
    f'"Operating EBIT share: NES proxy (enter per-sale EBIT share on Tab 3 col D for exact split).")',
    F(8,False,GREY,it=True)); hs.merge_cells("A94:J94")

# ==================== SAVE ====================
try: wb.calculation.fullCalcOnLoad=True; wb.calculation.calcId=0
except Exception as e: print("calc:",e)
wb.save(P)
print("SAVED:",P)
