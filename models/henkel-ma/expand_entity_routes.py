"""Expand entity routes from 2 to 8 slots.

Changes:
  Tab 3: rows 64-71 = 8 entity route rows (was 2 at 64-65)
         cols G-H added: spread override + computed spread
         row 72 = sum, row 73 = note
         old tax structure (69-72) cleared — inputs moved to E4

  E4:    B36 reads Tab 3 expanded range for weighted spread
         B94 now holds IC-lender jurisdiction directly (was Tab 3 B70)
         B96/B97 overrides moved to E4 C96/C97 (was Tab 3 B71/B72)
         B92/B93 rewritten to cover 8 entities via array formulas

  AT:    new check — entity shares sum to 100% (±1% tolerance)
"""

import openpyxl
from copy import copy

SRC = "/mnt/user-data/outputs/Henkel_MA_Financing_Model_teilverkauf_ccy.xlsx"
DST = SRC

wb = openpyxl.load_workbook(SRC)
ws_mix = wb["3. Scenario & Mix"]
ws_e4  = wb["E4 · KPIs & Funding"]
ws_at  = wb["✓ Acceptance Test"]

# ── helpers ─────────────────────────────────────────────────────────
def copy_style(src, dst):
    for attr in ("font", "alignment", "border", "fill", "number_format"):
        v = getattr(src, attr, None)
        if v:
            setattr(dst, attr, copy(v))

# Reference style from existing entity row 64
ref_cells = {c: ws_mix.cell(64, c) for c in range(1, 9)}

# ── 1. Save current tax-structure values before overwriting ────────
tax_lender = ws_mix.cell(70, 2).value  # IC-loan lender jurisdiction
tax_shield_override = ws_mix.cell(71, 2).value  # financing-shield tax override
wht_override = ws_mix.cell(72, 2).value  # WHT override

print(f"Tax inputs to migrate → lender={tax_lender!r}, shield={tax_shield_override!r}, wht={wht_override!r}")

# ── 2. Unmerge cells in rows 62-73 that would block writes ────────
to_unmerge = []
for mr in ws_mix.merged_cells.ranges:
    if mr.min_row >= 62 and mr.max_row <= 73:
        to_unmerge.append(str(mr))
for mr_str in to_unmerge:
    ws_mix.unmerge_cells(mr_str)
    print(f"  unmerged {mr_str}")

# Clear rows 66-73 on Tab 3 (old sum, note, empty, tax)
for r in range(66, 74):
    for c in range(1, 12):
        ws_mix.cell(r, c).value = None

# ── 3. Add column headers at row 63 ───────────────────────────────
ws_mix.cell(63, 7, "Spread override (bps)")
ws_mix.cell(63, 8, "Computed spread (bps)")
copy_style(ws_mix.cell(63, 6), ws_mix.cell(63, 7))
copy_style(ws_mix.cell(63, 6), ws_mix.cell(63, 8))

# ── 4. Write/update D and H formulas for all 8 entity rows ────────
for row in range(64, 72):
    # D: auto-derived PP share (expanded COUNTIF range)
    d_formula = (
        f'=IF(OR($A{row}="",$C{row}=""),0,'
        f'IFERROR(INDEX($C$8:$C$29,MATCH($C{row},$A$8:$A$29,0))'
        f'/COUNTIF($C$64:$C$71,$C{row}),0))'
    )
    ws_mix.cell(row, 4, d_formula)

    # H: computed spread (bps) — uses global params unless G overrides
    h_formula = (
        f'=IF($A{row}="","",'
        f'IF(ISNUMBER($G{row}),$G{row},'
        f'IF($E{row}="IC Loan",'
        f"'E4 · KPIs & Funding'!$B$24+'E4 · KPIs & Funding'!$B$25,"
        f"'E4 · KPIs & Funding'!$B$27)"
        f'+IF($F{row}="Y",'
        f"'E4 · KPIs & Funding'!$B$26,0)))"
    )
    ws_mix.cell(row, 8, h_formula)

    # Apply styling to new rows (66-71)
    if row >= 66:
        for c in range(1, 9):
            copy_style(ref_cells.get(c, ref_cells[1]), ws_mix.cell(row, c))

# ── 5. Σ sum at row 72 ────────────────────────────────────────────
ws_mix.cell(72, 1, "Σ PP share (auto-100% when ccy mix = 100%)")
ws_mix.cell(72, 4, "=SUM(D64:D71)")
ws_mix.cell(72, 8, "=IFERROR(SUMPRODUCT($D$64:$D$71,$H$64:$H$71)/SUM($D$64:$D$71),0)")
copy_style(ref_cells[1], ws_mix.cell(72, 1))

# ── 6. Note at row 73 ─────────────────────────────────────────────
ws_mix.cell(73, 1,
    "► Share auto-derives from currency mix; type a value in D to override.  "
    "Spread override (G) bypasses the global parameters on E4.  "
    "Blank rows are ignored.")

# ── 7. E4 B36: read weighted spread from Tab 3 expanded range ─────
ws_e4.cell(36, 2,
    "=IFERROR(SUMPRODUCT('3. Scenario & Mix'!$D$64:$D$71,"
    "'3. Scenario & Mix'!$H$64:$H$71)"
    "/SUM('3. Scenario & Mix'!$D$64:$D$71),0)")

# ── 8. Move tax inputs to E4 ──────────────────────────────────────
# B94: IC-loan lender jurisdiction (was Tab 3 B70)
ws_e4.cell(94, 2, tax_lender if tax_lender else "DE")
ws_e4.cell(94, 1, "IC-loan lender jurisdiction (edit here)")

# B96/B97: override mechanism — move to C96/C97 inputs
ws_e4.cell(96, 3, tax_shield_override if isinstance(tax_shield_override, (int, float)) else None)
ws_e4.cell(97, 3, wht_override if isinstance(wht_override, (int, float)) else None)

# B96: read override from C96, else use blended
ws_e4.cell(96, 2, "=IF(ISNUMBER($C$96),$C$96,$B$92)")
# B97: read override from C97, else use derived
ws_e4.cell(97, 2, "=IF(ISNUMBER($C$97),$C$97,$B$93)")

# D96/D97: explain where override lives
ws_e4.cell(96, 4,
    "type override in C96; blank = use blended CIT (was on tab 3)")
ws_e4.cell(97, 4,
    "type override in C97; blank = use derived WHT (was on tab 3)")

# ── 9. E4 B92/B93: blended CIT/WHT across 8 entities ─────────────
# B92: blended CIT — array SUMPRODUCT over Tab 3 entities
ws_e4.cell(92, 2,
    "=IFERROR(SUMPRODUCT("
    "'3. Scenario & Mix'!$D$64:$D$71,"
    "IFERROR(INDEX('Tax Reference'!$C$7:$C$21,"
    "MATCH('3. Scenario & Mix'!$B$64:$B$71,'Tax Reference'!$A$7:$A$21,0)),"
    "'1. Deal Inputs'!$B$42))"
    "/SUM('3. Scenario & Mix'!$D$64:$D$71),"
    "'1. Deal Inputs'!$B$42)")

# B93: effective WHT — only for IC Loan routes, using treaty matrix
ws_e4.cell(93, 2,
    "=IFERROR(SUMPRODUCT("
    "'3. Scenario & Mix'!$D$64:$D$71"
    "*('3. Scenario & Mix'!$E$64:$E$71=\"IC Loan\"),"
    "IFERROR(INDEX('Tax Reference'!$B$26:$P$40,"
    "MATCH($B$94,'Tax Reference'!$A$26:$A$40,0),"
    "MATCH('3. Scenario & Mix'!$B$64:$B$71,'Tax Reference'!$B$25:$P$25,0)),"
    "0)),0)")

# ── 10. Update E4 display mirror rows 31-32 (keep as first 2) ────
# They still reference 64-65, which is fine.  Add a note:
ws_e4.cell(33, 1, "Σ shares (display: first 2; full 8-entity blend at B36)")
# D33 stays as SUM(D31:D32) — it's just a display check for the first 2

# ── 11. AT tab: add entity-share sum check ─────────────────────────
# Find the input bounds section (row 89 = header, checks start at 90)
# Add after the last existing check (row 99)
at_row = 100
ws_at.cell(at_row, 1, "Entity PP-share sum ≈ 100% (within ±1%)")
ws_at.cell(at_row, 2,
    '=IF(ABS(SUM(\'3. Scenario & Mix\'!$D$64:$D$71)-1)>0.01,'
    '"❌ FAIL — entity shares sum to "&TEXT(SUM(\'3. Scenario & Mix\'!$D$64:$D$71),"0.0%")&" (not 100%)",'
    '"✅ PASS")')
ws_at.cell(at_row, 3,
    "Entities don't account for the full purchase price; check Tab 3 rows 64-71.")

# Style from existing check row
ref_at = ws_at.cell(90, 1)
for c in range(1, 4):
    copy_style(ws_at.cell(90, c), ws_at.cell(at_row, c))

# ── 12. Verify verdict range covers new check ─────────────────────
verdict = ws_at.cell(5, 2).value
assert "$B$120" in verdict, f"Verdict range doesn't reach row {at_row}: {verdict}"
assert at_row <= 120, f"Check row {at_row} > 120"

# ── SAVE ───────────────────────────────────────────────────────────
wb.save(DST)

print()
print("✅ Entity routes expanded from 2 → 8 slots")
print("   Tab 3 rows 64-71: entity data (A-H)")
print("   Tab 3 row 72: Σ sum")
print("   Tab 3 row 73: note")
print("   Tax inputs migrated to E4 (B94, C96, C97)")
print("   E4 B36: reads Tab 3 D64:D71 × H64:H71")
print("   E4 B92/B93: array formulas for 8-entity CIT/WHT blending")
print("   AT row 100: entity-share sum ≈ 100% check")
print(f"   Saved → {DST}")
