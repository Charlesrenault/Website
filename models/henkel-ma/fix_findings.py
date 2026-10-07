import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

P="/mnt/user-data/outputs/Henkel_MA_Financing_Model_teilverkauf_ccy.xlsx"
wb=openpyxl.load_workbook(P)

FONT="Arial"
def F(sz=10,b=False,color="000000",it=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=it)
def fill(c): return PatternFill("solid",fgColor=c)
YELLOW=fill("FFF2CC"); HEADERBG=fill("1F3864"); BANDBG=fill("EEF3FA")
thin=Side(style="thin",color="BFBFBF"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
BLUE="0000FF"; GREEN="008000"; GREY="555555"; WHITE="FFFFFF"; RED="CC0000"
PCT1='0.0%'
def put(ws,a,v,font=None,fillc=None,num=None,align=None,wrap=False,border=False):
    c=ws[a]; c.value=v; c.font=font or F()
    if fillc: c.fill=fillc
    if num: c.number_format=num
    if align or wrap: c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if border: c.border=BORDER
    return c

MIX="'3. Scenario & Mix'"; DI="'1. Deal Inputs'"
mix=wb["3. Scenario & Mix"]; hs=wb["Henkel Summary"]

# --- unmerge all ranges we need to rewrite ---
for ws in (mix, hs):
    for mr in list(ws.merged_cells.ranges):
        s=str(mr)
        # Tab3 row 92
        if ws==mix and mr.min_row==92 and mr.max_row==92:
            ws.unmerge_cells(s)
        # HS rows 61, 73, 82, 89, 90, 91, 94
        if ws==hs and mr.min_row in (61,73,82,89,90,91,94) and mr.max_row==mr.min_row:
            ws.unmerge_cells(s)

# ======================================================================
# F3: Sum(D) > 100% guard on Tab 3
# ======================================================================
put(mix,"A92","Each row removes that currency's financing leg. Col D = EBIT share (blank = NES proxy). Same-ccy fractions add up (cap 100%). Blank B = 100%.",F(8,False,GREY,it=True))
mix.merge_cells("A92:E92")
put(mix,"F92",
    '=IF(SUMPRODUCT(--ISNUMBER($D$86:$D$91),$D$86:$D$91)>1,'
    '"⛔ Σ EBIT shares > 100% ("&TEXT(SUMPRODUCT(--ISNUMBER($D$86:$D$91),$D$86:$D$91),"0%")&")","")',
    F(8,True,RED))
mix.merge_cells("F92:I92")

# F3 part 2: floor row 79 at 0 so negative retained is capped
for j in range(8):
    cl=get_column_letter(3+j)
    put(hs,f"{cl}79",f'=IFERROR(MAX(0,1-SUM({cl}120:{cl}125)),1)',F(9),None,PCT1,align="center")

# ======================================================================
# F1+F4: Echo table shows source + flags zero EBIT share
# ======================================================================
for k in range(6):
    r=83+k; tr=86+k; rr=96+k
    put(hs,f"E{r}",
        f'=IF({MIX}!$A${tr}="","",IF(AND($N${rr}=0,{MIX}!$A${tr}<>""),"0% ⚠",'
        f'IF(ISNUMBER({MIX}!$D${tr}),TEXT($N${rr},"0.0%")&" ✎",TEXT($N${rr},"0.0%")&" ~")))',
        F(8),None,None,align="center")

put(hs,"A82","Per sale: ccy · % leg · year · EBIT share (✎=input, ~=NES, ⚠=zero) · rate(Y1) · Σ interest",
    F(8,True,GREY))
hs.merge_cells("A82:J82")

# ======================================================================
# F2+F5: A94 note — per-sale granularity + B65="N" warning
# ======================================================================
put(hs,"A94",
    f'=IF({DI}!$B$65="N",'
    f'IF(SUMPRODUCT(--ISNUMBER({MIX}!$D$86:$D$91))>0,'
    f'"⚠ B65=N (funding-only): EBIT-share inputs in col D are IGNORED — operating side stays 100%. Set B65=Y for entity sales.",'
    f'"Funding-only sale (B65=N): no operating impact. Set B65=Y if selling an operating entity."),'
    f'IF(SUMPRODUCT(--ISNUMBER({MIX}!$D$86:$D$91))>0,'
    f'SUMPRODUCT(--ISNUMBER({MIX}!$D$86:$D$91))&" sale(s) with exact EBIT share (✎)"'
    f'&IF(SUMPRODUCT(--({MIX}!$A$86:$A$91<>""),--NOT(ISNUMBER({MIX}!$D$86:$D$91)))>0,'
    f'", "&SUMPRODUCT(--({MIX}!$A$86:$A$91<>""),--NOT(ISNUMBER({MIX}!$D$86:$D$91)))&" using NES proxy (~) — enter D for accuracy","")&". WHT scales with remaining interest.",'
    f'"Operating EBIT share: NES proxy for all sales (~). Enter per-sale EBIT share on Tab 3 col D for exact split."))',
    F(8,False,GREY,it=True))
hs.merge_cells("A94:J94")

# ======================================================================
# F7: Warning when blocks ④ + ⑤ both active
# ======================================================================
put(hs,"A73",
    f'=IF($L$73>0,'
    f'IF(AND(ISNUMBER({DI}!$B$57),{DI}!$B$57>0,ISNUMBER({DI}!$B$58)),'
    f'"⚠ BOTH blocks active: ④ (uniform "&TEXT({DI}!$B$57,"0%")&") AND ⑤ ("&$L$73&" targeted sale(s)). These are ADDITIVE on the financing side — check for double-counting.",'
    f'"Active: "&$L$73&" targeted sale(s) from the Tab 3 table — additive across currencies & years."),'
    f'"Inert — enter sales on Tab 3 (TARGETED SALES table). Block ④ above = uniform stake sale.")',
    F(8,False,GREY,it=True))
hs.merge_cells("A73:J73")

put(hs,"A61",
    f'=IF(AND(ISNUMBER({DI}!$B$57),{DI}!$B$57>0,ISNUMBER({DI}!$B$58),{DI}!$B$58>=1),'
    f'IF($L$73>0,'
    f'"⚠ BOTH blocks active: ④ (uniform "&TEXT({DI}!$B$57,"0%")&" from yr "&{DI}!$B$58&") AND ⑤ ("&$L$73&" targeted sale(s)). These are ADDITIVE — check for double-counting.",'
    f'"Active: "&TEXT({DI}!$B$57,"0%")&" sold from horizon year "&{DI}!$B$58&" → retained "&TEXT({DI}!$B$59,"0%")&". Amount, financing, income, EPS & NFP scale to retained share."),'
    f'"No partial sale configured — acquirer share = 100% (block inert). Activate on Deal Inputs B57/B58.")',
    F(8,False,GREY,it=True))
hs.merge_cells("A61:J61")

# ======================================================================
# Save
# ======================================================================
try: wb.calculation.fullCalcOnLoad=True; wb.calculation.calcId=0
except Exception as e: print("calc:",e)
wb.save(P)
print("SAVED with all fixes:", P)
