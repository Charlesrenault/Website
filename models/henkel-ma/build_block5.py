import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC="/tmp/claude-0/-home-user-Website/cdbc117f-3212-54dc-9478-5898308d5bbf/scratchpad/model.xlsx"
OUT="/tmp/claude-0/-home-user-Website/cdbc117f-3212-54dc-9478-5898308d5bbf/scratchpad/model_v2.xlsx"
wb=openpyxl.load_workbook(SRC)

FONT="Arial"
def F(sz=10,b=False,color="000000",it=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=it)
def fill(c): return PatternFill("solid",fgColor=c)
BLUE="0000FF"; GREEN="008000"; GREY="555555"; WHITE="FFFFFF"
YELLOW=fill("FFF2CC"); INPUTBG=fill("DDEBF7"); HEADERBG=fill("1F3864"); BANDBG=fill("EEF3FA"); OKBG=fill("E2EFDA")
thin=Side(style="thin",color="BFBFBF"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
NUM_MM='#,##0.0;(#,##0.0);"-"'; PCT1='0.0%'; PCT2='0.00%'
def put(ws,a,v,font=None,fillc=None,num=None,align=None,wrap=False,border=False):
    c=ws[a]; c.value=v; c.font=font or F()
    if fillc: c.fill=fillc
    if num: c.number_format=num
    if align or wrap: c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if border: c.border=BORDER
    return c

NY=8
E3="'E3 · Cost & Schedule'"; E1="'E1 · Rates & Spreads'"; DI="'1. Deal Inputs'"; MIX="'3. Scenario & Mix'"; D9="Dashboard!$D$9"
SOLD=f"{DI}!$B$62"; FRACc=f"{DI}!$B$63"; EVENT=f"{DI}!$B$58"; MF=f"{DI}!$B$24"
FRACU="$L$75"   # resolved fraction (default 100%)
I0="$L$73"; I0N="$L$74"; ACT="$L$72"; NES="$L$76"

# ---------------- Deal Inputs: add targeted-sale inputs (rows 62-64) ----------------
di=wb["1. Deal Inputs"]
put(di,"A62","► TARGETED SALE — sold currency (ISO, e.g. MXN)",F(10,True))
put(di,"B62",None,F(10,True,BLUE),YELLOW,None,align="center",border=True)
put(di,"D62","leave blank = no targeted-currency sale. Uses the SAME event year (B58) as above. For a UNIFORM stake sale use B57 instead (not both).",F(8,False,GREY,it=True)); di.merge_cells("D62:E62")
put(di,"A63","► ...fraction of that currency's financing leg sold (%)",F(10,True))
put(di,"B63",1.0,F(10,True,BLUE),YELLOW,PCT1,align="center",border=True)
put(di,"D63","100% = sell the whole currency leg (e.g. the entire Mexico/MXN financing).",F(8,False,GREY,it=True)); di.merge_cells("D63:E63")
put(di,"A64","► implied sold share of the financing book",F(10,True))
put(di,"B64",f'=IFERROR(IF(AND({SOLD}<>"",ISNUMBER(MATCH({SOLD},{E3}!$A$15:$A$36,0))),INDEX({E3}!$C$15:$C$36,MATCH({SOLD},{E3}!$A$15:$A$36,0))*IF(ISNUMBER({FRACc}),{FRACc},1),0),0)',F(10,False,GREEN),None,PCT1,align="center",border=True)
put(di,"D64","= (that currency's financing share) × (fraction sold). Drives the exact amount & cost removed.",F(8,False,GREY,it=True)); di.merge_cells("D64:E64")
dv=DataValidation(type="list",formula1='"ARS,AUD,CAD,CHF,CNY,DKK,EUR,GBP,HKD,HUF,IDR,INR,JPY,MXN,NZD,PLN,ROL,SEK,SGD,THB,TRY,USD"',allow_blank=True)
di.add_data_validation(dv); dv.add(di["B62"])

# ---------------- Henkel Summary: Block ⑤ (rows 72-86) ----------------
hs=wb["Henkel Summary"]
put(hs,"A72","⑤  TARGETED SALE BY CURRENCY  —  exact per-currency divestment (e.g. sell the MXN leg, not a uniform slice)",F(11,True,WHITE),HEADERBG); hs.merge_cells("A72:J72")
put(hs,"A73",f'=IF({ACT}=1,"Active: sell "&TEXT({FRACU},"0%")&" of the "&{SOLD}&" financing leg from horizon year "&{EVENT}&". The mix renormalises automatically; interest drops by exactly that leg (incl. its own rate — not a flat % of the book).","Inert — enter a sold currency on Deal Inputs B62 (+ fraction B63, event year B58).")',F(8,False,GREY,it=True)); hs.merge_cells("A73:J73")
# year header row 74 (C:J) + helpers in L
put(hs,"A74","EUR mm  /  rate",F(9,True))
for j in range(NY): put(hs,f"{get_column_letter(3+j)}74",f'=IFERROR({DI}!$B$39+{j},"")',F(9,True),BANDBG,'0',align="center",border=True)
# helpers (column L, outside the merges)
put(hs,"K74","helpers →",F(7,False,GREY,it=True),align="right")
put(hs,"L72",f'=IF(AND({SOLD}<>"",{I0}>0,ISNUMBER({EVENT})),1,0)',F(7,False,GREY))        # active flag
put(hs,"L73",f'=IFERROR(MATCH({SOLD},{E3}!$A$15:$A$36,0),0)',F(7,False,GREY))              # E3/E1 currency index
put(hs,"L74",f'=IFERROR(MATCH({SOLD},{MIX}!$A$8:$A$29,0),0)',F(7,False,GREY))              # tab3 currency index
put(hs,"L75",f'=IF(ISNUMBER({FRACc}),{FRACc},1)',F(7,False,GREY))                          # fraction used
put(hs,"L76",f'=IFERROR(INDEX({MIX}!$G$8:$G$29,{I0N})/SUM({MIX}!$G$8:$G$29),0)',F(7,False,GREY))  # sold ccy NES share

rowsE={75:"Sold-currency financing share (of book)",76:"Sold-currency all-in rate",
       77:"Sold-currency interest leg (EUR mm)",78:"► Post-sale interest — acquirer (EUR mm)",
       79:"   of which: share of TOTAL interest shed (%)",80:"Retained financing-amount factor",
       81:"Sold-currency share in the mix AFTER the sale",82:"Operating-earnings retained factor (NES-based)"}
for r,t in rowsE.items(): put(hs,f"A{r}",t,F(9,True if r==78 else False))
for j in range(NY):
    cl=get_column_letter(3+j); c116=get_column_letter(36+j); mf=f"*{MF}" if j==0 else ""
    fromev=f"AND({ACT}=1,{j}>={EVENT}-1)"
    put(hs,f"{cl}75",f'=IFERROR(INDEX({E3}!{cl}$15:{cl}$36,{I0}),0)',F(9),None,PCT1,align="center")
    rate=(f'IF({D9}="(instrument mix)",'
          f'({E3}!$K$42*INDEX({E1}!$AJ$10:$AJ$31,{I0})+{E3}!$K$43*INDEX({E1}!${c116}$10:${c116}$31,{I0})+{E3}!$K$44*INDEX({E1}!$AR$10:$AR$31,{I0}))/{E3}!$K$45,'
          f'IF({D9}="Short Fix",INDEX({E1}!$AJ$10:$AJ$31,{I0}),'
          f'IF({D9}="Short Float",INDEX({E1}!${c116}$10:${c116}$31,{I0}),INDEX({E1}!$AR$10:$AR$31,{I0}))))'
          f'+INDEX({E3}!$H$120:$H$141,{I0})+{E3}!$B$51/10000')
    put(hs,f"{cl}76",f'=IF({I0}=0,0,IFERROR({rate},0))',F(9),None,PCT2,align="center")
    put(hs,f"{cl}77",f'=IFERROR({cl}75*{E3}!{cl}$64*{cl}76{mf},0)',F(9),None,NUM_MM,align="center")
    put(hs,f"{cl}78",f'=IFERROR({E3}!{cl}$65-IF({fromev},{FRACU}*{cl}77,0),"")',F(9,True),None,NUM_MM,align="center")
    put(hs,f"{cl}79",f'=IFERROR(IF({E3}!{cl}$65>0.0001,IF({fromev},{FRACU}*{cl}77,0)/{E3}!{cl}$65,0),0)',F(9),None,PCT1,align="center")
    put(hs,f"{cl}80",f'=IFERROR(1-IF({fromev},{FRACU}*{cl}75,0),1)',F(9),None,PCT1,align="center")
    put(hs,f"{cl}81",f'=IFERROR(IF({fromev},{cl}75*(1-{FRACU})/(1-{FRACU}*{cl}75),{cl}75),0)',F(9),None,PCT1,align="center")
    put(hs,f"{cl}82",f'=IFERROR(1-IF({fromev},{FRACU}*{NES},0),1)',F(9),None,PCT1,align="center")
put(hs,"A83","Interest drops by the sold currency's OWN leg (share × balance × that currency's rate) — so selling a small, high-rate leg (e.g. MXN 5%) can cut far more than 5% of interest. The mix renormalises automatically.",F(8,False,GREY,it=True)); hs.merge_cells("A83:J83")
put(hs,"A84","► Post-sale NFP contribution (retained financing amount × Group NFP change)",F(9,True))
put(hs,"C84",f'=IFERROR($J$80*$G$10,"")',F(9,True,GREEN),None,NUM_MM,align="center",border=True)
put(hs,"D84","operating-earnings side uses the NES share (row 82); a precise entity-level P&L split is not in the single-blend plan (first-order).",F(8,False,GREY,it=True)); hs.merge_cells("D84:J84")
# cross-check cell: block⑤ rate for the sold ccy vs block② EUR/USD when sold ccy is EUR/USD
put(hs,"A85","Self-check: with sold ccy = EUR, row 76 must equal block ② EUR rate (row 17); with USD, row 20.",F(8,False,GREY,it=True)); hs.merge_cells("A85:J85")
hs.column_dimensions["L"].width=8; hs.column_dimensions["K"].width=9

wb.save(OUT)
print("SAVED",OUT)
print("sheets:",len(wb.sheetnames))
