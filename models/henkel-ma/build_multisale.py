import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

P="/mnt/user-data/outputs/Henkel_MA_Financing_Model_teilverkauf_ccy.xlsx"
wb=openpyxl.load_workbook(P)
FONT="Arial"
def F(sz=10,b=False,color="000000",it=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=it)
YELLOW=PatternFill("solid",fgColor="FFF2CC"); HEADERBG=PatternFill("solid",fgColor="1F3864"); BANDBG=PatternFill("solid",fgColor="EEF3FA")
thin=Side(style="thin",color="BFBFBF"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
BLUE="0000FF"; GREEN="008000"; GREY="555555"; WHITE="FFFFFF"
NUM_MM='#,##0.0;(#,##0.0);"-"'; PCT1='0.0%'; PCT2='0.00%'
def put(ws,a,v,font=None,fillc=None,num=None,align=None,wrap=False,border=False):
    c=ws[a]; c.value=v; c.font=font or F()
    if fillc: c.fill=fillc
    if num: c.number_format=num
    if align or wrap: c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if border: c.border=BORDER
    return c
NY=8
E3="'E3 · Cost & Schedule'"; E1="'E1 · Rates & Spreads'"; DI="'1. Deal Inputs'"; MIX="'3. Scenario & Mix'"; D9="Dashboard!$D$9"; MF=f"{DI}!$B$24"; ENT=f"{DI}!$B$65"
CCYS="ARS,AUD,CAD,CHF,CNY,DKK,EUR,GBP,HKD,HUF,IDR,INR,JPY,MXN,NZD,PLN,ROL,SEK,SGD,THB,TRY,USD"

di=wb["1. Deal Inputs"]; mix=wb["3. Scenario & Mix"]; hs=wb["Henkel Summary"]

# ---------- Deal Inputs: retire single-ccy inputs, point to the table ----------
di["A62"]="► Multiple TARGETED SALES → enter on Tab 3 (TARGETED SALES table)"; di["A62"].font=F(10,True)
di["B62"]=None; di["B63"]=None; di["B64"]=None
di["A63"]=None; di["A64"]=None
for a in ("D63","D64"):
    if di[a].value: di[a]=None
di["D62"]="currency · % of leg · event year, one row per sale (up to 6). B65 below applies to all."; di["D62"].font=F(8,False,GREY,it=True)
try: di.merge_cells("D62:E62")
except Exception: pass

# ---------- Tab 3: 6-row TARGETED SALES table (rows 84-92) ----------
put(mix,"A84","TARGETED SALES BY CURRENCY  (feeds Henkel Summary ⑤) — up to 6; any currencies, same or different years; additive",F(11,True,WHITE),HEADERBG); mix.merge_cells("A84:I84")
put(mix,"A85","Sold currency (ISO)",F(8,True),BANDBG,align="center",wrap=True,border=True)
put(mix,"B85","% of that currency's leg",F(8,True),BANDBG,align="center",wrap=True,border=True)
put(mix,"C85","Event year (1–8)",F(8,True),BANDBG,align="center",wrap=True,border=True)
put(mix,"D85","(leave a row blank = no sale)",F(8,False,GREY,it=True))
for k in range(6):
    r=86+k
    put(mix,f"A{r}",None,F(9,True,BLUE),YELLOW,None,align="center",border=True)
    put(mix,f"B{r}",None,F(9,True,BLUE),YELLOW,PCT1,align="center",border=True)
    put(mix,f"C{r}",None,F(9,True,BLUE),YELLOW,'0',align="center",border=True)
put(mix,"A92","Each row removes that currency's financing leg (amount + its own rate) from its event year. Fractions of the SAME currency add up (cap at 100%). Blank % = 100% of the leg.",F(8,False,GREY,it=True)); mix.merge_cells("A92:I92")
dvc=DataValidation(type="list",formula1=f'"{CCYS}"',allow_blank=True); mix.add_data_validation(dvc)
for k in range(6): dvc.add(mix[f"A{86+k}"])

# ---------- Henkel Summary: clear old single-sale block ⑤ body (73-90) + helpers L72:L76 ----------
# unmerge any range overlapping rows 72-91 (old block ⑤ notes) so we can rewrite freely
for mr in list(hs.merged_cells.ranges):
    if mr.max_row>=72 and mr.min_row<=91:
        hs.unmerge_cells(str(mr))
for r in range(72,92):
    for col in "ABCDEFGHIJKLMNOP":
        c=hs[f"{col}{r}"]
        try:
            if c.value is not None: c.value=None
        except Exception: pass

# ---------- hidden helper matrices (rows 96-125) ----------
# per sale k: row 96+k => C:J rate, L i0, M i0nes, N nes ; shed_int 104+k ; shed_amt 112+k ; shed_nes 120+k
for k in range(6):
    tr=86+k                      # tab3 sales row
    ccyK=f"{MIX}!$A${tr}"; fracK=f"IF(ISNUMBER({MIX}!$B${tr}),{MIX}!$B${tr},1)"; yrK=f"{MIX}!$C${tr}"
    rr=96+k
    put(hs,f"L{rr}",f'=IFERROR(MATCH({ccyK},{E3}!$A$15:$A$36,0),0)',F(7,False,GREY))
    put(hs,f"M{rr}",f'=IFERROR(MATCH({ccyK},{MIX}!$A$8:$A$29,0),0)',F(7,False,GREY))
    put(hs,f"N{rr}",f'=IFERROR(INDEX({MIX}!$G$8:$G$29,$M${rr})/SUM({MIX}!$G$8:$G$29),0)',F(7,False,GREY))
    i0=f"$L${rr}"
    for j in range(NY):
        cl=get_column_letter(3+j); c116=get_column_letter(36+j); mfj=f"*{MF}" if j==0 else ""
        gated=f'AND({ccyK}<>"",{i0}>0,ISNUMBER({yrK}),{j}>={yrK}-1)'
        rate=(f'IF({D9}="(instrument mix)",'
              f'({E3}!$K$42*INDEX({E1}!$AJ$10:$AJ$31,{i0})+{E3}!$K$43*INDEX({E1}!${c116}$10:${c116}$31,{i0})+{E3}!$K$44*INDEX({E1}!$AR$10:$AR$31,{i0}))/{E3}!$K$45,'
              f'IF({D9}="Short Fix",INDEX({E1}!$AJ$10:$AJ$31,{i0}),'
              f'IF({D9}="Short Float",INDEX({E1}!${c116}$10:${c116}$31,{i0}),INDEX({E1}!$AR$10:$AR$31,{i0}))))'
              f'+INDEX({E3}!$H$120:$H$141,{i0})+{E3}!$B$51/10000')
        shareK=f'INDEX({E3}!{cl}$15:{cl}$36,{i0})'
        put(hs,f"{cl}{rr}",      f'=IF({i0}=0,0,IFERROR({rate},0))',F(7),None,PCT2)                                   # rate_k(j)
        put(hs,f"{cl}{104+k}",   f'=IFERROR(IF({gated},{fracK}*{shareK}*{E3}!{cl}$64*{cl}{rr}{mfj},0),0)',F(7),None,NUM_MM)  # shed interest
        put(hs,f"{cl}{112+k}",   f'=IFERROR(IF({gated},{fracK}*{shareK},0),0)',F(7),None,PCT1)                        # shed amount share
        put(hs,f"{cl}{120+k}",   f'=IFERROR(IF(AND({gated},{ENT}<>"N"),{fracK}*$N${rr},0),0)',F(7),None,PCT1)        # shed NES (operating)
for r in range(95,127): hs.row_dimensions[r].hidden=True
for col in ("L","M","N"): hs.column_dimensions[col].width=7

# ---------- main Block ⑤ (aggregate) ----------
put(hs,"A72","⑤  TARGETED SALES BY CURRENCY  —  per-currency divestments (table on Tab 3) · additive, multiple currencies & years",F(11,True,WHITE),HEADERBG); hs.merge_cells("A72:J72")
put(hs,"L73",f'=SUMPRODUCT(--({MIX}!$A$86:$A$91<>""),--ISNUMBER({MIX}!$C$86:$C$91))',F(7,False,GREY))   # #active sales
put(hs,"A73",f'=IF($L$73>0,"Active: "&$L$73&" targeted sale(s) from the Tab 3 table — additive across currencies & years.","Inert — enter sales on Tab 3 (TARGETED SALES table). Block ④ above = uniform stake sale.")',F(8,False,GREY,it=True)); hs.merge_cells("A73:J73")
put(hs,"A74","EUR mm  /  rate",F(9,True))
for j in range(NY): put(hs,f"{get_column_letter(3+j)}74",f'=IFERROR({DI}!$B$39+{j},"")',F(9,True),BANDBG,'0',align="center",border=True)
rowsE={75:"► Total interest shed — all sales (EUR mm)",76:"► Post-sale interest — acquirer (EUR mm)",
       77:"   share of TOTAL interest shed (%)",78:"Retained financing-amount factor",
       79:"Operating-earnings retained factor (entity sales; NES)",80:"► Post-sale NFP contribution (retained amount × Group NFP)"}
for r,t in rowsE.items(): put(hs,f"A{r}",t,F(9,True if r in(76,) else False))
for j in range(NY):
    cl=get_column_letter(3+j)
    put(hs,f"{cl}75",f'=IFERROR(SUM({cl}104:{cl}109),0)',F(9),None,NUM_MM,align="center")
    put(hs,f"{cl}76",f'=IFERROR({E3}!{cl}$65-{cl}75,"")',F(9,True),None,NUM_MM,align="center")
    put(hs,f"{cl}77",f'=IFERROR(IF({E3}!{cl}$65>0.0001,{cl}75/{E3}!{cl}$65,0),0)',F(9),None,PCT1,align="center")
    put(hs,f"{cl}78",f'=IFERROR(1-SUM({cl}112:{cl}117),1)',F(9),None,PCT1,align="center")
    put(hs,f"{cl}79",f'=IFERROR(1-SUM({cl}120:{cl}125),1)',F(9),None,PCT1,align="center")
    put(hs,f"{cl}80",f'=IFERROR({cl}78*$G$10,"")',F(9),None,NUM_MM,align="center")
# per-sale echo table
put(hs,"A82","Per sale (echo): currency · % · year · rate(Y1) · Σ interest shed (horizon)",F(8,True,GREY)); hs.merge_cells("A82:J82")
for k in range(6):
    r=83+k; tr=86+k; rr=96+k
    put(hs,f"A{r}",f'="Sale "&{k+1}',F(8))
    put(hs,f"B{r}",f'=IFERROR({MIX}!$A${tr},"")',F(8,False,GREEN),None,align="center")
    put(hs,f"C{r}",f'=IF({MIX}!$A${tr}="","",IF(ISNUMBER({MIX}!$B${tr}),{MIX}!$B${tr},1))',F(8),None,PCT1,align="center")
    put(hs,f"D{r}",f'=IFERROR({MIX}!$C${tr},"")',F(8),None,'0',align="center")
    put(hs,f"E{r}",f'=IF({MIX}!$A${tr}="","",$C${rr})',F(8),None,PCT2,align="center")
    put(hs,f"F{r}",f'=IF({MIX}!$A${tr}="","",SUM(C{104+k}:J{104+k}))',F(8),None,NUM_MM,align="center")
put(hs,"A89","Rows 75–80 aggregate ALL sales (additive). Per-sale echo: E = that currency's Y1 rate (if EUR, must equal Block ② EUR rate, row 17 — self-check); F = total interest it removes over the horizon.",F(8,False,GREY,it=True)); hs.merge_cells("A89:J89")
put(hs,"A90","Same-currency rows add up (cap 100% of that leg). Operating side via NES (entity toggle B65). NFP per year = retained amount × Group NFP; sale proceeds (cash-in) still a memo.",F(8,False,GREY,it=True)); hs.merge_cells("A90:J90")

try: wb.calculation.fullCalcOnLoad=True; wb.calculation.calcId=0
except Exception as e: print("calc:",e)
wb.save(P)
print("SAVED multi-sale version:",P)
print("sheets:",len(wb.sheetnames))
